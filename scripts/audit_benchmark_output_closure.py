#!/usr/bin/env python3
"""Fail closed when a benchmark freezes before official outputs and pilots are closed."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime
from pathlib import Path


def finding(severity: str, code: str, message: str, evidence=None) -> dict:
    item = {"severity": severity, "code": code, "message": message}
    if evidence is not None:
        item["evidence"] = evidence
    return item


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def timestamp(value: object) -> datetime:
    if not value:
        raise ValueError("timestamp is empty")
    parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        raise ValueError("timestamp has no timezone")
    return parsed


def resolve(root: Path, value: object) -> Path:
    candidate = Path(str(value or ""))
    return candidate if candidate.is_absolute() else root / candidate


def json_pointer(payload: object, pointer: str) -> object:
    if pointer == "":
        return payload
    if not pointer.startswith("/"):
        raise ValueError("JSON pointer must start with /")
    current = payload
    for raw in pointer[1:].split("/"):
        token = raw.replace("~1", "/").replace("~0", "~")
        if isinstance(current, list):
            current = current[int(token)]
        elif isinstance(current, dict):
            current = current[token]
        else:
            raise KeyError(token)
    return current


def nonempty(value: object) -> bool:
    return value is not None and value != "" and value != [] and value != {}


def run_content_check(path: Path, check: dict) -> tuple[bool, object]:
    kind = check.get("type")
    if kind == "file_nonempty":
        minimum = int(check.get("minimum_bytes", 1))
        size = path.stat().st_size
        return size >= minimum, {"size": size, "minimum_bytes": minimum}
    if kind == "json_pointer_nonempty":
        pointer = str(check.get("pointer", ""))
        payload = json.loads(path.read_text(encoding="utf-8"))
        value = json_pointer(payload, pointer)
        return nonempty(value), {"pointer": pointer, "value_type": type(value).__name__}
    if kind == "text_marker":
        marker = str(check.get("marker", ""))
        present = bool(marker) and marker in path.read_text(encoding="utf-8")
        return present, {"marker": marker, "present": present}
    if kind == "xlsx_range_nonempty":
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as exc:
            raise RuntimeError("openpyxl is required for xlsx_range_nonempty") from exc
        sheet = str(check.get("sheet", ""))
        cell_range = str(check.get("range", ""))
        minimum = int(check.get("minimum_nonempty", 1))
        book = load_workbook(path, read_only=True, data_only=False)
        if sheet not in book.sheetnames:
            return False, {"sheet": sheet, "available": book.sheetnames}
        cells = book[sheet][cell_range]
        count = sum(cell.value is not None and cell.value != "" for row in cells for cell in row)
        return count >= minimum, {"sheet": sheet, "range": cell_range, "nonempty": count, "minimum": minimum}
    raise ValueError(f"unsupported content check: {kind!r}")


def inspect(payload: dict, root: Path, schema: Path | None = None) -> dict:
    findings: list[dict] = []
    if schema:
        try:
            import jsonschema
            jsonschema.validate(payload, json.loads(schema.read_text(encoding="utf-8")))
        except ModuleNotFoundError:
            findings.append(finding("Major", "JSONSCHEMA_UNAVAILABLE", "Closure manifest was not checked with JSON Schema."))
        except Exception as exc:
            findings.append(finding("Blocker", "CLOSURE_SCHEMA", str(exc)))

    pre = payload.get("pre_pilot_manifest") or {}
    pre_path = resolve(root, pre.get("path"))
    if not pre_path.is_file():
        findings.append(finding("Blocker", "PRE_PILOT_MANIFEST_MISSING", "Pre-pilot manifest is missing.", str(pre_path)))
    else:
        expected = str(pre.get("sha256") or "")
        actual = sha256(pre_path)
        if actual != expected:
            findings.append(finding("Blocker", "PRE_PILOT_HASH_MISMATCH", "Pre-pilot manifest hash does not match.", {"expected": expected, "actual": actual}))
    try:
        frozen_at = timestamp(pre.get("frozen_at"))
        started_at = timestamp(pre.get("first_pilot_started_at"))
        if frozen_at > started_at:
            findings.append(finding("Blocker", "PILOT_FREEZE_ORDER", "Pre-pilot manifest was frozen after pilot execution began."))
    except Exception as exc:
        findings.append(finding("Blocker", "PILOT_TIMESTAMP_INVALID", str(exc)))

    closed = 0
    requirements = payload.get("requirements") or []
    for requirement in requirements:
        requirement_id = str(requirement.get("id", "<missing>"))
        if requirement.get("required") and requirement.get("status") != "complete":
            findings.append(finding("Blocker", "OFFICIAL_OUTPUT_INCOMPLETE", f"Required output {requirement_id} is not complete.", requirement.get("artifact_path")))
            continue
        if requirement.get("status") != "complete":
            continue
        path = resolve(root, requirement.get("artifact_path"))
        if not path.is_file():
            findings.append(finding("Blocker", "OUTPUT_ARTIFACT_MISSING", f"Output artifact for {requirement_id} is missing.", str(path)))
            continue
        actual = sha256(path)
        expected = str(requirement.get("artifact_sha256") or "")
        if actual != expected:
            findings.append(finding("Blocker", "OUTPUT_HASH_MISMATCH", f"Output artifact hash mismatch for {requirement_id}.", {"path": str(path), "expected": expected, "actual": actual}))
            continue
        checks = requirement.get("content_checks") or []
        if not checks:
            findings.append(finding("Blocker", "OUTPUT_CONTENT_CHECK_MISSING", f"Complete output {requirement_id} has no content check."))
            continue
        failed = False
        for check in checks:
            try:
                passed, evidence = run_content_check(path, check)
            except Exception as exc:
                passed, evidence = False, str(exc)
            if not passed:
                failed = True
                findings.append(finding("Blocker", "OUTPUT_CONTENT_CHECK_FAILED", f"Content check failed for {requirement_id}.", evidence))
        if not failed:
            closed += 1

    blockers = [item for item in findings if item["severity"] in {"Blocker", "Critical"}]
    return {
        "case_id": payload.get("case_id"),
        "status": "pass" if not blockers else "blocked",
        "requirement_count": len(requirements),
        "closed_requirement_count": closed,
        "findings": findings,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--root", type=Path)
    parser.add_argument("--schema", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--strict", action="store_true")
    args = parser.parse_args()
    payload = json.loads(args.manifest.read_text(encoding="utf-8"))
    root = (args.root or args.manifest.parent).resolve()
    schema = args.schema or Path(__file__).resolve().parent.parent / "assets" / "schemas" / "benchmark-output-closure.schema.json"
    report = inspect(payload, root, schema)
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    severities = {item["severity"] for item in report["findings"]}
    return int(bool({"Blocker", "Critical"} & severities or (args.strict and "Major" in severities)))


if __name__ == "__main__":
    raise SystemExit(main())
